from __future__ import annotations

import io
import os
import re
from html import escape

from aiogram import Bot, F, Router
from aiogram.fsm.context import FSMContext
from aiogram.types import BufferedInputFile, CallbackQuery, FSInputFile, Message
from aiogram.utils.keyboard import InlineKeyboardBuilder
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import app_state
import backup
import config
from callbacks import AdmCB, StatsCB
from database import crud, queries
from database.models import async_session
from handlers import views
from keyboards.admin_kb import admin_main_kb, cancel_kb
from states import AdminManagerFSM

router = Router(name="admin_tools")
router.message.filter(F.from_user.id.in_(config.ADMIN_IDS))
router.callback_query.filter(F.from_user.id.in_(config.ADMIN_IDS))


@router.callback_query(StatsCB.filter())
async def show_stats(callback: CallbackQuery, callback_data: StatsCB) -> None:
    async with async_session() as session:
        text = await queries.stats_text(session, callback_data.period)
    kb = InlineKeyboardBuilder()
    for key, label in queries.PERIODS.items():
        mark = "• " if key == callback_data.period else ""
        kb.button(text=f"{mark}{label}", callback_data=StatsCB(period=key))
    kb.adjust(2)
    foot = InlineKeyboardBuilder()
    foot.button(text="🔙 Адмін-меню", callback_data=AdmCB(section="home"))
    kb.attach(foot)
    await views.reply_or_edit(callback, text, kb.as_markup())
    await callback.answer()


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, header in enumerate(headers, 1):
        longest = max([len(str(header))] + [len(str(r[i - 1])) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), 60)


@router.callback_query(AdmCB.filter(F.section == "export"))
async def export_excel(callback: CallbackQuery) -> None:
    async with async_session() as session:
        clients, orders = await queries.export_rows(session)
    wb = Workbook()
    _sheet(wb, "Клієнти", queries.CLIENT_HEADERS, clients, first=True)
    _sheet(wb, "Замовлення", queries.ORDER_HEADERS, orders)
    buffer = io.BytesIO()
    wb.save(buffer)
    await callback.message.answer_document(
        BufferedInputFile(buffer.getvalue(), filename="jobin_clients.xlsx"),
        caption=f"📥 Клієнтів: {len(clients)} · замовлень: {len(orders)}",
        reply_markup=admin_main_kb(),
    )
    await callback.answer()


@router.callback_query(AdmCB.filter(F.section == "backup"))
async def manual_backup(callback: CallbackQuery) -> None:
    path, name = await backup.create_backup()
    try:
        await callback.message.answer_document(
            FSInputFile(path, filename=name), caption="💾 Резервна копія бази. Зберігайте її в надійному місці.",
            reply_markup=admin_main_kb(),
        )
    finally:
        os.remove(path)
    await callback.answer()


def _link(code: str) -> str:
    return f"https://t.me/{app_state.bot_username}?start=m_{code}"


@router.callback_query(AdmCB.filter((F.section == "mgr") & (F.action == "list")))
async def mgr_list(callback: CallbackQuery, state: FSMContext) -> None:
    await state.clear()
    async with async_session() as session:
        managers = await crud.get_managers(session, only_active=False)
    lines = ["👥 <b>Менеджери</b>", "Кожен менеджер дає клієнтам своє посилання — так видно, хто привів клієнта.", ""]
    kb = InlineKeyboardBuilder()
    for m in managers:
        mark = "" if m.is_active else " 🔴"
        lines.append(f"<b>{escape(m.name)}</b>{mark} · @{escape(m.username)}\n<code>{_link(m.code)}</code>\n")
        kb.button(text=f"✏️ {m.name}", callback_data=AdmCB(section="mgr", action="view", id=m.id))
    kb.button(text="➕ Додати менеджера", callback_data=AdmCB(section="mgr", action="add"))
    kb.button(text="🔙 Адмін-меню", callback_data=AdmCB(section="home"))
    kb.adjust(1)
    await views.reply_or_edit(callback, "\n".join(lines), kb.as_markup())
    await callback.answer()


async def _mgr_card(callback_or_message, manager_id: int) -> None:
    async with async_session() as session:
        m = await crud.get_manager(session, manager_id)
    kb = InlineKeyboardBuilder()
    for action, label in (("e_name", "Ім'я"), ("e_username", "Юзернейм"), ("e_code", "Код посилання")):
        kb.button(text=f"✏️ {label}", callback_data=AdmCB(section="mgr", action=action, id=manager_id))
    kb.button(
        text="🔴 Вимкнути" if m.is_active else "🟢 Увімкнути",
        callback_data=AdmCB(section="mgr", action="toggle", id=manager_id),
    )
    kb.button(text="🔙 До списку", callback_data=AdmCB(section="mgr", action="list"))
    kb.adjust(1)
    text = (
        f"<b>{escape(m.name)}</b> · @{escape(m.username)}\n"
        f"Код: <code>{escape(m.code)}</code>\nСтатус: {'активний' if m.is_active else 'вимкнений'}\n\n"
        f"Посилання для клієнтів:\n<code>{_link(m.code)}</code>"
    )
    await views.reply_or_edit(callback_or_message, text, kb.as_markup())


@router.callback_query(AdmCB.filter((F.section == "mgr") & (F.action == "view")))
async def mgr_view(callback: CallbackQuery, callback_data: AdmCB) -> None:
    await _mgr_card(callback, callback_data.id)
    await callback.answer()


@router.callback_query(AdmCB.filter((F.section == "mgr") & (F.action == "toggle")))
async def mgr_toggle(callback: CallbackQuery, callback_data: AdmCB) -> None:
    async with async_session() as session:
        m = await crud.get_manager(session, callback_data.id)
        await crud.update_manager_field(session, m.id, "is_active", not m.is_active)
    await _mgr_card(callback, callback_data.id)
    await callback.answer("Змінено")


EDIT_FIELDS = {
    "e_name": ("name", "нове ім'я"),
    "e_username": ("username", "новий Telegram-юзернейм (без @)"),
    "e_code": ("code", "новий код посилання (латиницею, цифри, _)"),
}


@router.callback_query(AdmCB.filter((F.section == "mgr") & F.action.in_(set(EDIT_FIELDS))))
async def mgr_edit_start(callback: CallbackQuery, callback_data: AdmCB, state: FSMContext) -> None:
    field, label = EDIT_FIELDS[callback_data.action]
    await state.clear()
    await state.update_data(manager_id=callback_data.id, field=field)
    await state.set_state(AdminManagerFSM.waiting_value)
    await callback.message.answer(f"Введіть {label}:", reply_markup=cancel_kb())
    await callback.answer()


def _valid_code(code: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z0-9_]{2,32}", code))


@router.message(AdminManagerFSM.waiting_value, F.text)
async def mgr_edit_finish(message: Message, state: FSMContext) -> None:
    data = await state.get_data()
    value = message.text.strip().lstrip("@")
    if data["field"] == "code" and not _valid_code(value):
        await message.answer("Код: 2–32 символи, лише латиниця, цифри та _. Спробуйте ще раз:", reply_markup=cancel_kb())
        return
    if not value:
        await message.answer("Значення не може бути порожнім:", reply_markup=cancel_kb())
        return
    try:
        async with async_session() as session:
            await crud.update_manager_field(session, data["manager_id"], data["field"], value)
    except Exception:
        await message.answer("Не вдалося зберегти (можливо, такий код уже існує). Спробуйте інший:", reply_markup=cancel_kb())
        return
    await state.clear()
    await _mgr_card(message, data["manager_id"])


@router.callback_query(AdmCB.filter((F.section == "mgr") & (F.action == "add")))
async def mgr_add_start(callback: CallbackQuery, state: FSMContext) -> None:
    await state.clear()
    await state.set_state(AdminManagerFSM.name)
    await callback.message.answer("Введіть ім'я менеджера (як його бачитимете в статистиці):", reply_markup=cancel_kb())
    await callback.answer()


@router.message(AdminManagerFSM.name, F.text)
async def mgr_add_name(message: Message, state: FSMContext) -> None:
    await state.update_data(name=message.text.strip())
    await state.set_state(AdminManagerFSM.username)
    await message.answer("Введіть Telegram-юзернейм менеджера (без @):", reply_markup=cancel_kb())


@router.message(AdminManagerFSM.username, F.text)
async def mgr_add_username(message: Message, state: FSMContext) -> None:
    await state.update_data(username=message.text.strip().lstrip("@"))
    await state.set_state(AdminManagerFSM.code)
    await message.answer("Придумайте короткий код для посилання (латиницею, наприклад ivan):", reply_markup=cancel_kb())


@router.message(AdminManagerFSM.code, F.text)
async def mgr_add_code(message: Message, state: FSMContext) -> None:
    code = message.text.strip()
    if not _valid_code(code):
        await message.answer("Код: 2–32 символи, лише латиниця, цифри та _. Спробуйте ще раз:", reply_markup=cancel_kb())
        return
    data = await state.get_data()
    try:
        async with async_session() as session:
            manager = await crud.create_manager(session, data["name"], data["username"], code)
    except Exception:
        await message.answer("Такий код уже існує. Введіть інший:", reply_markup=cancel_kb())
        return
    await state.clear()
    await _mgr_card(message, manager.id)